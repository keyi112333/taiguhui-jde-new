#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Excel读写处理类
"""
from pathlib import Path
from typing import List, Dict, Any, Optional
from log.loggingConfig import logger

try:
    import openpyxl
    from openpyxl import Workbook
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False
    logger.warning("需要安装 openpyxl: pip install openpyxl")


class ExcelHandler:
    """Excel读写处理类

    提供Excel文件的读取、写入、修改等操作
    """

    def __init__(self, file_path: str):
        self.file_path = Path(file_path)
        self.workbook: Optional[Workbook] = None
        self._load()

    def _load(self) -> None:
        """加载Excel文件"""
        if not HAS_OPENPYXL:
            raise ImportError("需要安装 openpyxl: pip install openpyxl")

        if self.file_path.exists():
            logger.info(f"加载Excel文件: {self.file_path}")
            self.workbook = openpyxl.load_workbook(self.file_path)
        else:
            logger.info(f"创建新Excel文件: {self.file_path}")
            self.workbook = Workbook()
            # 删除默认的Sheet
            if "Sheet" in self.workbook.sheetnames:
                self.workbook.remove(self.workbook["Sheet"])

    def read_sheet(self, sheet_name: str, header_row: int = 1) -> List[Dict[str, Any]]:
        """读取工作表数据为字典列表

        Args:
            sheet_name: 工作表名称
            header_row: 标题行号（默认为1）

        Returns:
            List[Dict]: 数据列表，每行数据为字典
        """
        if sheet_name not in self.workbook.sheetnames:
            logger.error(f"工作表不存在: {sheet_name}")
            return []

        sheet = self.workbook[sheet_name]
        logger.info(f"读取工作表: {sheet_name}")

        # 获取标题行
        headers = []
        for cell in sheet.iter_rows(min_row=header_row, max_row=header_row, values_only=True):
            headers = [str(h) if h else f"col_{i}" for i, h in enumerate(cell)]

        # 读取数据行
        data = []
        for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
            row_dict = {}
            for i, value in enumerate(row):
                if i < len(headers):
                    row_dict[headers[i]] = value
            data.append(row_dict)

        logger.info(f"读取数据行数: {len(data)}")
        return data

    def write_sheet(self, sheet_name: str, data: List[Dict[str, Any]], headers: Optional[List[str]] = None) -> None:
        """写入数据到工作表

        Args:
            sheet_name: 工作表名称
            data: 数据列表
            headers: 标题列表（可选，默认使用第一行数据的键）
        """
        # 创建或获取工作表
        if sheet_name in self.workbook.sheetnames:
            sheet = self.workbook[sheet_name]
            sheet.delete_rows(1, sheet.max_row)  # 清空现有数据
        else:
            sheet = self.workbook.create_sheet(sheet_name)

        # 确定标题
        if headers is None and data:
            headers = list(data[0].keys())

        if headers:
            # 写入标题行
            for col, header in enumerate(headers, start=1):
                sheet.cell(row=1, column=col, value=header)

            # 写入数据行
            for row_idx, row_data in enumerate(data, start=2):
                for col, header in enumerate(headers, start=1):
                    sheet.cell(row=row_idx, column=col, value=row_data.get(header))

        logger.info(f"写入数据行数: {len(data)}")

    def get_sheet_names(self) -> List[str]:
        """获取所有工作表名称"""
        return self.workbook.sheetnames

    def create_sheet(self, sheet_name: str) -> None:
        """创建新工作表"""
        if sheet_name not in self.workbook.sheetnames:
            self.workbook.create_sheet(sheet_name)
            logger.info(f"创建工作表: {sheet_name}")

    def delete_sheet(self, sheet_name: str) -> None:
        """删除工作表"""
        if sheet_name in self.workbook.sheetnames:
            self.workbook.remove(self.workbook[sheet_name])
            logger.info(f"删除工作表: {sheet_name}")

    def save(self, path: Optional[str] = None) -> None:
        """保存Excel文件

        Args:
            path: 保存路径（可选，默认使用原路径）
        """
        save_path = Path(path) if path else self.file_path
        self.workbook.save(save_path)
        logger.info(f"Excel文件已保存: {save_path}")

    def close(self) -> None:
        """关闭Excel文件"""
        if self.workbook:
            self.workbook.close()
            logger.info("Excel文件已关闭")

    def get_cell_value(self, sheet_name: str, row: int, col: int) -> Any:
        """获取单元格值"""
        sheet = self.workbook[sheet_name]
        return sheet.cell(row=row, column=col).value

    def set_cell_value(self, sheet_name: str, row: int, col: int, value: Any) -> None:
        """设置单元格值"""
        sheet = self.workbook[sheet_name]
        sheet.cell(row=row, column=col, value=value)

    def get_row_count(self, sheet_name: str) -> int:
        """获取行数"""
        sheet = self.workbook[sheet_name]
        return sheet.max_row

    def get_col_count(self, sheet_name: str) -> int:
        """获取列数"""
        sheet = self.workbook[sheet_name]
        return sheet.max_column