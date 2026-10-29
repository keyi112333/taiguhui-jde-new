"""从租房产税底稿更新：金额使用 Decimal 汇总，保留 MOH 公式行。"""
from collections import defaultdict
from copy import copy
from decimal import Decimal, InvalidOperation
from pathlib import Path
from tempfile import NamedTemporaryFile
import re
import zipfile

from openpyxl.formula.translate import Translator
from config.config_loader import ConfigError

EXTERNAL_LINK_DIR = "xl/externalLinks/"


def external_link_parts(path):
    """读取外部链接部件；openpyxl 保存只写回部分关系，须在保存后回写。"""
    with zipfile.ZipFile(path) as archive:
        return {name: archive.read(name) for name in archive.namelist()
                if name.startswith(EXTERNAL_LINK_DIR) and not name.endswith("/")}


def restore_external_links(path, parts):
    """把原始外部链接部件写回，避免 rId 悬空导致办公软件无法打开生成文件。"""
    if not parts:
        return
    path = Path(path)
    temp = None
    try:
        with zipfile.ZipFile(path) as archive:
            entries = [(item, archive.read(item.filename)) for item in archive.infolist()]
        with NamedTemporaryFile(dir=path.parent, suffix=".xlsx", delete=False) as handle:
            temp = Path(handle.name)
        with zipfile.ZipFile(temp, "w") as archive:
            for item, data in entries:
                archive.writestr(item, parts.get(item.filename, data))
        temp.replace(path)
        temp = None
    finally:
        if temp and temp.exists():
            temp.unlink()


def text(value):
    return str(value).strip() if value is not None else ""


def update_workbook(book, include_children=False):
    """新版模板以 Check 筛选条件为输入；旧版缺少条件时明确拒绝。"""
    from workflows.jde_flow.check_update import update_from_check
    return update_from_check(book, include_children)


def update_legacy(book, include_children=False):
    audit = book["房产税（从租）审核底稿"]
    ledger = book["JDE明细账"]
    calc = book["房产税（从租）计算底稿"]
    def columns(sheet):
        return {text(c.value): c.column for c in sheet[1] if c.value is not None}
    ac, lc = columns(audit), columns(ledger)
    def merged_value(row, col):
        for area in audit.merged_cells.ranges:
            if area.min_row <= row <= area.max_row and area.min_col <= col <= area.max_col:
                return audit.cell(area.min_row, area.min_col).value
        return audit.cell(row, col).value
    accounts = {}
    bu_order = []
    for row in range(2, audit.max_row + 1):
        account = text(audit.cell(row, ac["收入科目"]).value)
        if not account:
            continue
        bu = text(merged_value(row, ac["BU"]))
        if not bu:
            raise ConfigError(f"审核底稿第 {row} 行收入科目缺少 BU")
        if bu == "250108":
            # 固定从 MOH 取数，不生成来自明细账的重复 BU 行。
            continue
        if account in accounts and accounts[account] != bu:
            raise ConfigError(f"收入科目对应多个 BU：{account}")
        accounts[account] = bu
        if bu not in bu_order:
            bu_order.append(bu)
    if not accounts:
        raise ConfigError("审核底稿没有收入科目")
    details = defaultdict(Decimal)
    descriptions = {}
    for row in ledger.iter_rows(min_row=2, values_only=True):
        account = text(row[lc["帐号"] - 1])
        matches = [key for key in accounts if account == key or
                   (include_children and account.startswith(key + "."))]
        if not matches:
            continue
        bus = {accounts[key] for key in matches}
        if len(bus) != 1:
            raise ConfigError(f"科目匹配多个 BU：{account}")
        bu = bus.pop()
        doc = text(row[lc["单据 类型"] - 1])
        raw = row[lc["分类帐类型 1 金额"] - 1]
        if raw is None or text(raw) == "":
            raise ConfigError(f"科目 {account} 的金额为空")
        try:
            amount = Decimal(str(raw))
            if not amount.is_finite():
                raise InvalidOperation
        except (InvalidOperation, ValueError):
            raise ConfigError(f"科目 {account} 的金额不是有效数值：{raw}")
        if not doc:
            raise ConfigError(f"科目 {account} 的单据类型为空")
        details[bu, doc, account] += amount
        descriptions[account] = row[lc["帐户 说明"] - 1]
    if not details:
        raise ConfigError("审核底稿收入科目未匹配到明细账，未更新模板")

    start = 9  # 示例布局 B7:D8 为表头，数据自第 9 行起。
    total_rows = [r for r in range(start, calc.max_row + 1) if text(calc.cell(r, 2).value) == "合计"]
    if len(total_rows) != 1:
        raise ConfigError("计算底稿必须包含唯一的 B 列‘合计’行")
    old_total = total_rows[0]
    fixed = []
    old_keys, seen = {}, defaultdict(int)
    for row in range(start, old_total):
        bu, value = text(calc.cell(row, 2).value), calc.cell(row, 3).value
        if not bu:
            continue
        if bu == "250108":
            fixed.append((row, [copy(calc.cell(row, c)) for c in (2, 3, 4)]))
        else:
            old_keys[row] = (bu, seen[bu])
            seen[bu] += 1
    if len(fixed) > 1:
        raise ConfigError("计算底稿的固定 BU 250108 只能有一行")
    if not fixed:
        cells = [copy(calc.cell(start, c)) for c in (2, 3, 4)]
        cells[0].value = 250108
        cells[1].value = "=LOOKUP($C$4,MOH!$E$6:$AF$6,MOH!$E$7:$AF$7)"
        fixed.append((None, cells))
    elif fixed[0][1][1].data_type != "f":
        fixed[0][1][1].value = "=LOOKUP($C$4,MOH!$E$6:$AF$6,MOH!$E$7:$AF$7)"
    # 每个收入科目、单据类型单独成行；BU 相同也不能再次合并。
    entries = [(bu, doc, account, -details[bu, doc, account]) for bu in bu_order
               for account in sorted({a for b, d, a in details if b == bu})
               for doc in sorted((d for b, d, a in details if b == bu and a == account), reverse=True)]
    new_total = start + len(entries) + len(fixed) + 1
    style = [copy(calc.cell(start, c)._style) for c in (2, 3, 4)]
    total_style = [copy(calc.cell(old_total, c)._style) for c in (2, 3, 4)]
    # 仅移动 B:D 区域，避免移动旁边独立的分录区和 Check 区。
    for row in range(start, max(old_total, new_total) + 1):
        for col in (2, 3, 4):
            calc.cell(row, col).value = None
    new_bu_rows, new_keys, seq = defaultdict(list), {}, defaultdict(int)
    for row, (bu, doc, account, amount) in enumerate(entries, start):
        for col, fmt in zip((2, 3, 4), style):
            calc.cell(row, col)._style = copy(fmt)
        calc.cell(row, 2).value = int(bu) if bu.isdigit() else bu
        calc.cell(row, 3).value = float(amount)
        calc.cell(row, 4).value = f"=ROUND(C{row}*$C$8,2)"
        new_bu_rows[bu].append(row)
        new_keys[bu, seq[bu]] = row
        seq[bu] += 1
    relocation = {old: new_keys.get(key) for old, key in old_keys.items()}
    for row, (old, cells) in enumerate(fixed, start + len(entries)):
        for col, original in zip((2, 3, 4), cells):
            cell = calc.cell(row, col)
            cell._style = copy(original._style)
            formula = (re.sub(r"(?<![\w!])C4\b", "$C$4", original.value)
                       if original.data_type == "f" else original.value)
            cell.value = (Translator(formula, origin=original.coordinate).translate_formula(cell.coordinate)
                          if original.data_type == "f" else original.value)
        # 税额列始终按当前行收入 × 固定税率计算，保留两位小数。
        calc.cell(row, 4).value = f"=ROUND(C{row}*$C$8,2)"
        if old is not None:
            relocation[old] = row
        new_bu_rows[text(cells[0].value)].append(row)
    relocation[old_total] = new_total
    calc.cell(new_total, 2).value = "合计"
    for col in (3, 4):
        calc.cell(new_total, col).value = f"=SUM({chr(64+col)}{start}:{chr(64+col)}{new_total-1})"
    for col, fmt in zip((2, 3, 4), total_style):
        calc.cell(new_total, col)._style = fmt
    # 借项按分录帐号中的 BU 引用全部对应税额行。
    for row in range(8, calc.max_row + 1):
        account = text(calc.cell(row, 7).value)
        if account.endswith(".917013"):
            refs = new_bu_rows.get(account.split(".")[0], [])
            calc.cell(row, 8).value = "=" + "+".join(f"$D${r}" for r in refs) if refs else "=0"
        cell = calc.cell(row, 9)
        if cell.data_type == "f":
            def replace(match):
                old = int(match[2])
                if old not in relocation:
                    return match[0]
                new = relocation[old]
                if new is None:
                    # 对应 BU/单据类型无记录时不再生成该行，引用金额归零。
                    return "0"
                return match[1] + str(new)
            cell.value = re.sub(r"(\$?D\$?)(\d+)\b", replace, cell.value)
    # Check 明细保留科目与单据类型，便于追溯 BU 汇总。
    for row in range(5, max(calc.max_row, 5 + len(details)) + 1):
        for col in range(13, 17):
            calc.cell(row, col).value = None
    for row, ((bu, doc, account), amount) in enumerate(sorted(details.items()), 5):
        for col, value in enumerate((doc, descriptions[account], account, float(amount)), 13):
            calc.cell(row, col).value = value
    calc.cell(5 + len(details), 16).value = f"=SUM(P5:P{4+len(details)})"
    return [{"BU": bu, "收入科目": account, "单据类型": doc,
             "汇总金额": str(-amount), "计税金额": str(amount)}
            for bu, doc, account, amount in entries]
