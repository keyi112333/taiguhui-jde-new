"""运行结果表写入与企业微信机器人通知。"""
import re
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from log.loggingConfig import logger


RESULT_FILE_TEMPLATE = "房产税计税表更新与分录录入运行结果表-{year}.xlsx"
RESULT_HEADERS = (
    "会计期间", "公司", "公司代码", "分录批号",
    "运行结果", "异常情况", "运行开始时间", "运行结束时间",
)
RESULT_WIDTHS = (12, 16, 12, 20, 14, 40, 20, 20)
ROBOT_SHEET_PATTERN = re.compile(r"^WechatRobot(\d+)$", re.IGNORECASE)
WEBHOOK_KEYS = ("webhook", "webhook_url", "url", "地址", "机器人地址")


def result_file_path(config, year=None):
    """运行结果表路径：日志目录下按年份保存一份。"""
    year = year or datetime.now().year
    directory = config.get_path("log")
    if not isinstance(directory, (str, Path)):
        raise TypeError(f"日志目录配置无效：{directory!r}")
    return Path(directory) / RESULT_FILE_TEMPLATE.format(year=year)


def _style_header(sheet):
    for cell in sheet[1]:
        cell.font = Font(name="宋体", size=11)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for index, width in enumerate(RESULT_WIDTHS, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"


def append_run_result(config, record, year=None):
    """把一次运行记录追加到年度运行结果表，返回文件路径。"""
    path = result_file_path(config, year)
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.is_file():
        book = load_workbook(path)
        sheet = book.active
        if sheet.max_row == 1:
            _style_header(sheet)
    else:
        book = Workbook()
        sheet = book.active
        sheet.title = "运行结果"
        sheet.append(list(RESULT_HEADERS))
        _style_header(sheet)
    values = [record.get(header) for header in RESULT_HEADERS]
    sheet.append(["" if value is None else value for value in values])
    row = sheet.max_row
    sheet.cell(row, 6).alignment = Alignment(vertical="top", wrap_text=True)
    for column in (7, 8):
        sheet.cell(row, column).number_format = "yyyy/m/d h:mm:ss"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(RESULT_HEADERS))}{row}"
    book.save(path)
    book.close()
    logger.info(f"运行结果已记录：{path}（{record.get('公司') or ''} {record.get('运行结果') or ''}）")
    return path


def build_notification_message(company, period, total, success, failed, filename):
    """按约定格式生成企业微信通知文本。"""
    return (
        f"各位好，请查收 {company}_房产税计税表更新与分录录入，谢谢！\n"
        f"本次处理会计期间：{period}\n"
        f"总处理行数：{total} 行，处理成功：{success} 个，"
        f"失败/待补运行：{failed} 个，具体查看《{filename}》"
    )


def _robot_sheets(config):
    try:
        names = list(config.sheet_names())
    except (AttributeError, TypeError):
        return []
    sheets = []
    for name in names:
        match = ROBOT_SHEET_PATTERN.fullmatch(str(name).strip())
        if match:
            sheets.append((int(match.group(1)), name))
    return [name for _, name in sorted(sheets)]


def _read_webhook(config, sheet):
    try:
        rows = config.get_rows(sheet)
    except (AttributeError, TypeError):
        rows = []
    for row in rows:
        values = [row.get(key) for key in WEBHOOK_KEYS] + list(row.values())
        for value in values:
            if isinstance(value, str) and value.strip().lower().startswith("http"):
                return value.strip()
    for key in WEBHOOK_KEYS:
        try:
            value = config.get(f"{sheet}.{key}", "")
        except (AttributeError, TypeError):
            value = ""
        if isinstance(value, str) and value.strip().lower().startswith("http"):
            return value.strip()
    return ""


def _check_response(response):
    response.raise_for_status()
    payload = response.json()
    if payload.get("errcode", 0) != 0:
        raise RuntimeError(f"企业微信返回错误：{payload}")
    return payload


def _send_webhook(webhook, payload, timeout=15):
    import requests

    return _check_response(requests.post(webhook, json=payload, timeout=timeout))


def _upload_file(webhook, path, timeout=60):
    import requests

    with Path(path).open("rb") as handle:
        response = requests.post(
            webhook,
            params={"type": "file"},
            files={"media": (Path(path).name, handle, "application/octet-stream")},
            timeout=timeout,
        )
    payload = _check_response(response)
    media_id = payload.get("media_id")
    if not media_id:
        raise RuntimeError(f"企业微信上传文件未返回 media_id：{payload}")
    return media_id


def notify_wechat_robots(config, message, file_path=None, timeout=15):
    """逐个机器人发送通知；未配置 Webhook 的机器人跳过，单个失败不影响其他机器人。"""
    sent = 0
    for sheet in _robot_sheets(config):
        webhook = _read_webhook(config, sheet)
        if not webhook:
            logger.info(f"{sheet} 未配置 Webhook，跳过通知")
            continue
        try:
            _send_webhook(
                webhook, {"msgtype": "text", "text": {"content": message}}, timeout
            )
            if file_path and Path(file_path).is_file():
                media_id = _upload_file(webhook, file_path)
                _send_webhook(
                    webhook, {"msgtype": "file", "file": {"media_id": media_id}}, timeout
                )
            sent += 1
            logger.info(f"{sheet} 通知已发送")
        except Exception as exc:
            logger.warning(f"{sheet} 通知发送失败：{exc}")
    if sent == 0:
        logger.info("没有可用的企业微信机器人，未发送通知")
    return sent
