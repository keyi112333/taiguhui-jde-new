"""从 R55GL1005 的 Cat23 分组构建审核底稿。"""
from copy import copy
from openpyxl.workbook.defined_name import DefinedName
from config.config_loader import ConfigError


def norm(value):
    if isinstance(value, (float, int)) and value == int(value):
        return str(int(value))
    return str(value).strip() if value is not None else ''


def update_audit(book, calc_total):
    source = book['R55GL1005']
    target = book['房产税（从租）审核底稿']
    headers = [
        r 
        for r in range(1, min(source.max_row, 30) + 1)       
        if any(
            norm(source.cell(r, c).value) == 'Cat23'
            for c in range(1, min(source.max_column, 10) + 1)
        )
    ]
    if len(headers) != 1:
        raise ConfigError('R55GL1005 无法唯一定位 Cat23 表头')
    header = headers[0]
    expected = ['Cat23', 'Object Account', 'Subsidiary', 'Subledger', 'Account Description']
    if [norm(source.cell(header, c).value) for c in range(1, 6)] != expected:
        raise ConfigError('R55GL1005 表头结构与示例不符')
    if [norm(source.cell(header, c).value) for c in (7, 8)] != ['Debits', 'Credits']:
        raise ConfigError('R55GL1005 缺少本期借贷列')
    # 使用完整科目标识跟踪分类和样式，避免插入行后蓝色区域错位。
    known = {}
    for r in range(8, target.max_row + 1):
        if target.cell(r, 6).value is not None:
            key = tuple(norm(target.cell(r, c).value) for c in (3, 4, 5, 6))
            known[key] = [copy(target.cell(r, c)._style) for c in range(2, 10)]
    selected = []
    active = False
    groups = {}
    for r in range(header + 1, source.max_row + 1):
        cat = norm(source.cell(r, 1).value)
        desc = norm(source.cell(r, 5).value)
        if cat:
            active = desc in ('主营业务收入', '其他业务收入')
            if active:
                if desc in groups:
                    raise ConfigError(f'R55GL1005 收入分组重复：{desc}')
                groups[desc] = 8 + len(selected)
        if active and desc:
            values = [source.cell(r, c).value for c in (1, 2, 3, 4, 5, 7, 8)]
            key = tuple(norm(v) for v in values[1:5])
            if key not in known:
                raise ConfigError(f'审核底稿出现未分类科目，请先复核计税分类：{key}')
            selected.append((values, known[key]))
    if set(groups) != {'主营业务收入', '其他业务收入'}:
        raise ConfigError('R55GL1005 缺少主营业务收入或其他业务收入分组')
    for r in range(8, max(target.max_row, 7 + len(selected)) + 1):
        for c in range(2, 10):
            target.cell(r, c).value = None
    blue = []
    for r, (values, styles) in enumerate(selected, 8):
        for c, value in enumerate(values, 2):
            target.cell(r, c).value = value
        for c, style in enumerate(styles, 2):
            target.cell(r, c)._style = copy(style)
        target.cell(r, 9).value = f'=H{r}-G{r}'
        color = target.cell(r, 9).fill.fgColor
        if color.type == 'rgb' and color.rgb == 'FF00B0F0':
            blue.append(r)
    if not blue:
        raise ConfigError('审核底稿缺少不计税科目分类')
    book.defined_names.add(DefinedName('蓝色', attr_text=','.join(
        f"'房产税（从租）审核底稿'!$I${r}" for r in blue)))
    for r, desc in ((8, '主营业务收入'), (9, '其他业务收入')):
        target.cell(r, 13).value = f'=F{groups[desc]}'
        target.cell(r, 14).value = f'=I{groups[desc]}'
    target['N10'] = '=SUM(蓝色)'
    target['N11'] = '=N8+N9-N10'
    target['N12'] = '=LOOKUP($D$4,MOH!$E$6:$AF$6,MOH!$E$7:$AF$7)'
    target['N14'] = f"=N11+N12-'房产税（从租）计算底稿'!C{calc_total}"
    target['D4'] = "='房产税（从租）计算底稿'!C4"
