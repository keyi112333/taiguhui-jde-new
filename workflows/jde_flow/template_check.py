"""用办公软件重新计算后核对实际单元格值，不使用历史公式缓存。"""
from decimal import Decimal, InvalidOperation
import os
from pathlib import Path
import re
import uuid

from config.config_loader import ConfigError
from log.loggingConfig import logger


def recalculate_file(path, engine="Excel.Application"):
    """启动独立实例；关闭外链更新和宏，仅保存重算后的生成文件。"""
    import zipfile
    import openpyxl
    import pythoncom
    import win32com.client

    path = Path(path).resolve()
    if not path.is_file():
        raise ConfigError(f"公式重算文件不存在：{path}")
    if path.stat().st_size == 0:
        raise ConfigError(f"公式重算文件为空：{path}")
    try:
        with zipfile.ZipFile(path) as archive:
            damaged_member = archive.testzip()
        if damaged_member:
            raise ConfigError(f"公式重算文件损坏，XLSX 成员校验失败：{damaged_member}")
        check_book = openpyxl.load_workbook(path, read_only=True, data_only=False)
        check_book.close()
    except ConfigError:
        raise
    except (OSError, zipfile.BadZipFile, KeyError, ValueError) as exc:
        raise ConfigError(f"公式重算文件不是有效的 XLSX 工作簿：{path}；{exc}") from exc

    pythoncom.CoInitialize()
    app = book = None
    try:
        app = win32com.client.DispatchEx(engine)
        app.Visible = False
        app.DisplayAlerts = False
        app.AutomationSecurity = 3
        try:
            book = app.Workbooks.Open(
                Filename=str(path),
                UpdateLinks=0,
                ReadOnly=False,
                IgnoreReadOnlyRecommended=True,
                AddToMru=False,
            )
        except Exception as exc:
            raise ConfigError(
                f"Excel 无法打开公式重算文件：{path} "
                f"（{path.stat().st_size} 字节）。请确认文件未被其他 Excel 实例占用，"
                "并尝试用桌面 Excel 手动打开该文件。"
            ) from exc
        if book.ReadOnly:
            raise ConfigError("生成文件被占用，无法保存公式重算结果")
        app.CalculateFullRebuild()
        book.Save()
    finally:
        try:
            if book is not None:
                book.Close(SaveChanges=False)
        finally:
            try:
                if app is not None:
                    app.Quit()
            finally:
                pythoncom.CoUninitialize()


INVALID_SHEET_CHARS = re.compile(r"[\\/*?:\[\]]")


def _normalize_used_values(values, rows, cols):
    """把 COM 的 UsedRange.Value 统一成 rows×cols 的二维元组。"""
    if rows <= 0 or cols <= 0:
        return ()
    if rows == 1 and cols == 1:
        return ((values,),)
    if rows == 1:
        return (tuple(values),)
    if cols == 1:
        return tuple((item,) for item in values)
    return tuple(tuple(row) for row in values)


def _trim_empty_tail(values):
    """去掉末尾的全空行与全空列，避免把宏撑大的空白区域写进结果。"""
    rows = [list(row) for row in values]
    while rows and not any(value is not None for value in rows[-1]):
        rows.pop()
    width = max((len(row) for row in rows), default=0)
    while width:
        column = [row[width - 1] if len(row) >= width else None for row in rows]
        if any(value is not None for value in column):
            break
        width -= 1
    return tuple(tuple(row[:width]) for row in rows)


def _safe_sheet_title(name, used_titles):
    title = INVALID_SHEET_CHARS.sub("_", str(name or "Sheet"))[:31] or "Sheet"
    base = title
    index = 2
    while title in used_titles:
        suffix = f"_{index}"
        title = base[:31 - len(suffix)] + suffix
        index += 1
    used_titles.add(title)
    return title


def _write_sheets_to_xlsx(sheets, destination):
    """按原行列位置把读取到的值写入新的 .xlsx，不保留公式与格式。"""
    import openpyxl

    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    used_titles = set()
    for name, values, first_row, first_col in sheets:
        sheet = workbook.create_sheet(
            title=_safe_sheet_title(name, used_titles)
        )
        for row_index, row in enumerate(values, start=first_row):
            for col_index, value in enumerate(row, start=first_col):
                sheet.cell(row=row_index, column=col_index, value=value)
    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)
    workbook.close()


def _trusted_excel_dirs():
    """读取 Excel 受信任位置；只有受信任位置里的文件才会执行模板宏。"""
    import winreg

    found = []
    try:
        with winreg.OpenKey(winreg.HKEY_CURRENT_USER,
                            r"Software\Microsoft\Office") as office:
            versions = [
                winreg.EnumKey(office, index)
                for index in range(winreg.QueryInfoKey(office)[0])
            ]
    except OSError:
        return found

    for version in versions:
        key_path = (
            rf"Software\Microsoft\Office\{version}"
            r"\Excel\Security\Trusted Locations"
        )
        try:
            with winreg.OpenKey(winreg.HKEY_CURRENT_USER, key_path) as locations:
                count = winreg.QueryInfoKey(locations)[0]
                for index in range(count):
                    name = winreg.EnumKey(locations, index)
                    try:
                        with winreg.OpenKey(locations, name) as location:
                            path, _ = winreg.QueryValueEx(location, "Path")
                    except OSError:
                        continue
                    found.append(path)
        except OSError:
            continue
    return found


def _writable_trusted_dir():
    """返回可写的受信任目录；跳过 XLSTART，避免 Excel 启动时自动加载临时文件。"""
    for raw in _trusted_excel_dirs():
        directory = Path(os.path.expandvars(raw))
        if not directory.is_dir():
            continue
        if "xlstart" in str(directory).lower():
            continue
        try:
            probe = directory / f".jde_write_probe_{uuid.uuid4().hex}"
            probe.write_text("", encoding="utf-8")
            probe.unlink()
        except OSError:
            continue
        return directory
    return None


def convert_legacy_excel_to_xlsx(source, destination, engine="Excel.Application"):
    """把 JDE 导出的旧版 .xls 转为 openpyxl 可读的 .xlsx。

    JDE 的输出模板自带宏，Excel 打开时会提示“宏已被禁用/启用内容”。
    要让宏真正运行（等同点击“启用内容”），文件必须位于 Excel 受信任位置，
    因此先把源文件复制到可写的受信任目录再打开；打开后优先让 Excel 直接
    另存为 .xlsx（保留格式），只有另存被阻止时才退回读取单元格值重建。
    """
    import pythoncom
    import shutil
    import win32com.client

    source = Path(source).resolve()
    destination = Path(destination).resolve()
    if source == destination:
        raise ConfigError("R55GL1005 转换源文件与目标文件不能相同")
    if not source.is_file():
        raise ConfigError(f"R55GL1005 源文件不存在：{source}")
    destination.parent.mkdir(parents=True, exist_ok=True)

    trusted_dir = _writable_trusted_dir()
    working_source = source
    temp_copy = None
    temp_xlsx = None
    if trusted_dir is None:
        logger.warning(
            "未找到可写的 Excel 受信任位置，模板宏可能不会运行；"
            "请在信任中心添加受信任位置以得到与“启用内容”一致的结果"
        )
    else:
        token = uuid.uuid4().hex
        temp_copy = trusted_dir / f"_jde_r55_{token}.xls"
        temp_xlsx = trusted_dir / f"_jde_r55_{token}.xlsx"
        shutil.copy2(source, temp_copy)
        working_source = temp_copy
        logger.info(f"R55GL1005 已复制到受信任位置以运行模板宏：{trusted_dir}")

    sheets = []
    saved = False
    pythoncom.CoInitialize()
    app = book = None
    try:
        app = win32com.client.DispatchEx(engine)
        app.Visible = False
        app.DisplayAlerts = False
        app.AutomationSecurity = 1  # 启用内容：允许 JDE 模板宏运行
        # 不能关闭 EnableEvents：JDE 模板的处理逻辑挂在 Workbook_Open 事件上，
        # 关闭事件会让宏不执行，读到的仍是未处理的数据。
        book = app.Workbooks.Open(
            str(working_source), UpdateLinks=0, ReadOnly=False
        )
        save_target = temp_xlsx if temp_xlsx is not None else destination
        try:
            book.SaveAs(str(save_target), 51)
            saved = save_target.is_file() and save_target.stat().st_size > 0
            if not saved:
                logger.warning(
                    "Excel 另存为 xlsx 未生成文件（通常是被安全策略阻止），"
                    "改为读取单元格值重建"
                )
        except Exception as exc:
            logger.warning(f"Excel 另存为 xlsx 失败，改为读取单元格值重建：{exc}")
            saved = False

        if not saved:
            logger.info(
                f"R55GL1005 已启用模板宏打开：工作表数={book.Sheets.Count}"
            )
            for index in range(1, book.Sheets.Count + 1):
                sheet = book.Sheets(index)
                used = sheet.UsedRange
                logger.info(
                    f"  工作表 {sheet.Name}：{used.Rows.Count} 行 × "
                    f"{used.Columns.Count} 列（起始 "
                    f"{used.Row},{used.Column}）"
                )
                values = _trim_empty_tail(
                    _normalize_used_values(
                        used.Value, used.Rows.Count, used.Columns.Count
                    )
                )
                sheets.append((sheet.Name, values, used.Row, used.Column))
    finally:
        try:
            if book is not None:
                book.Close(SaveChanges=False)
        finally:
            try:
                if app is not None:
                    app.Quit()
            finally:
                pythoncom.CoUninitialize()
                if temp_xlsx is not None and saved:
                    try:
                        shutil.move(str(temp_xlsx), str(destination))
                    except OSError as exc:
                        logger.warning(f"移动转换结果失败：{exc}")
                        saved = False
                for leftover in (temp_copy, temp_xlsx):
                    if leftover is None:
                        continue
                    try:
                        if leftover.exists():
                            leftover.unlink()
                    except OSError as exc:
                        logger.warning(f"删除受信任位置临时文件失败：{exc}")

    if saved:
        logger.info("R55GL1005 已由 Excel 直接另存为 xlsx（保留原格式）")
        return
    if not sheets:
        raise ConfigError(f"R55GL1005 未读取到任何工作表：{source}")
    _write_sheets_to_xlsx(sheets, destination)
    if not destination.is_file() or destination.stat().st_size == 0:
        raise ConfigError(f"R55GL1005 转换后文件无效：{destination}")


def check_sheet(sheet, tolerance="0.00", audit=None):
    """定位表头、合计、分录和 Check 记录，不依赖固定数据行数。"""
    def label(value):
        return re.sub(r"\s+", "", str(value)) if value is not None else ""
    def number(row, col):
        cell = sheet.cell(row, col)
        try:
            if cell.value is None or isinstance(cell.value, bool):
                raise ValueError
            value = Decimal(str(cell.value))
            if not value.is_finite():
                raise ValueError
            return value
        except (InvalidOperation, ValueError):
            raise ConfigError(f"核对无法取数：{sheet.title}!{cell.coordinate}={cell.value!r}，请确认公式重算成功")
    try:
        limit = Decimal(str(tolerance))
        if not limit.is_finite() or limit < 0:
            raise ValueError
    except (InvalidOperation, ValueError):
        raise ConfigError("核对容差必须为非负金额")
    def unique(cells, description):
        if len(cells) != 1:
            raise ConfigError(f"无法唯一定位{description}：找到 {len(cells)} 处")
        return cells[0]
    bu_header = unique([c for row in sheet for c in row if label(c.value) == "BU"], "BU 表头")
    bu_col = bu_header.column
    total = unique([sheet.cell(r, bu_col) for r in range(bu_header.row + 1, sheet.max_row + 1)
                    if label(sheet.cell(r, bu_col).value) == "合计"], "收入/税额合计行")
    income, tax = number(total.row, bu_col + 1), number(total.row, bu_col + 2)
    fixed = unique([sheet.cell(r, bu_col) for r in range(bu_header.row + 1, total.row)
                    if label(sheet.cell(r, bu_col).value) == "250108"], "固定 BU 250108 行")
    number(fixed.row, bu_col + 1)
    debit_header = unique([c for row in sheet for c in row if label(c.value) == "借项金额"], "借项金额表头")
    debit_rows = [r for r in range(debit_header.row + 1, sheet.max_row + 1)
                  if label(sheet.cell(r, debit_header.column - 1).value).endswith(".917013")]
    if not debit_rows:
        raise ConfigError("没有找到借项分录行")
    debit = sum((number(r, debit_header.column) for r in debit_rows), Decimal(0))
    check_header = unique([c for row in sheet for c in row
                           if label(c.value) == "分类帐类型1金额"], "Check 金额表头")
    check_rows = [r for r in range(check_header.row + 1, sheet.max_row + 1)
                  if label(sheet.cell(r, check_header.column - 1).value)]
    if not check_rows:
        raise ConfigError("Check 区没有科目明细")
    check_total_row = unique([sheet.cell(r, check_header.column) for r in
                             range(max(check_rows) + 1, sheet.max_row + 1)
                             if sheet.cell(r, check_header.column).value is not None], "Check 合计")
    check_total = number(check_total_row.row, check_total_row.column)
    check_details = sum((number(r, check_header.column) for r in check_rows), Decimal(0))
    def result(name, left, right, locations):
        delta = left - right
        return {"核对项": name, "左值": str(left), "右值": str(right),
                "差额": str(delta), "容差": str(limit), "通过": abs(delta) <= limit,
                "位置": locations}
    if audit is not None:
        original = sheet
        try:
            sheet = audit
            taxable, moh = number(11, 14), number(12, 14)
            difference = number(14, 14)
        finally:
            sheet = original
        return [result("税额合计=借项汇总", tax, debit,
                       {"税额合计": sheet.cell(total.row, bu_col+2).coordinate, "借项行": debit_rows}),
                result("收入合计=审核计税收入+MOH", income, taxable + moh,
                       {"审核计税收入": audit.title + "!N11", "MOH": audit.title + "!N12"}),
                result("审核差额=0", difference, Decimal(0), {"差额": audit.title + "!N14"}),
                result("Check合计=Check明细汇总", check_total, check_details, {"Check明细行": check_rows})]
    return [result("税额合计=借项汇总", tax, debit,
                   {"税额合计": sheet.cell(total.row, bu_col+2).coordinate, "借项行": debit_rows}),
            result("收入合计=Check取反", income, -check_total,
                   {"收入合计": sheet.cell(total.row, bu_col+1).coordinate,
                    "固定BU": fixed.coordinate, "Check合计": check_total_row.coordinate}),
            result("Check合计=Check明细汇总", check_total, check_details,
                   {"Check明细行": check_rows})]
