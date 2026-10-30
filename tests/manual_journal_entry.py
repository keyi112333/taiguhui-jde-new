#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""模块六联调（非离线）：真实登录 JDE，录入日记帐分录并回写批号。

真实执行内容：打开浏览器 → 登录 JDE → 打开“日记帐分录” → 粘贴计算底稿 G-J →
停在填写完成的页面（不点取消、不保存），等待人工核对后按回车结束。

源工作簿不会被改动：先复制到输出目录，批号只写入副本。

运行：
    python tests/manual_journal_entry.py
    python tests/manual_journal_entry.py --workbook "data/房产税计税表 太古汇 202609更新.xlsx"
"""
import argparse
import shutil
from datetime import date, datetime
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import openpyxl

from config import Config
from core.context import TaskContext
from log.loggingConfig import logger
from web.browser import BrowserManager
from web.JDE.jde_function import create_journal_entry
from workflows.jde_flow.login_task import JDELoginTask
from workflows.jde_flow.property_tax import PropertyTaxFlow, format_jde_date


DEFAULT_WORKBOOK = ROOT / "data/房产税计税表 太古汇 202609更新.xlsx"
DEFAULT_OUTPUT = ROOT / "data/JDE流程/03_输出"
SHEET_NAME = "房产税（从租）计算底稿"


def cell_text(value):
    if isinstance(value, float):
        return format(value, "f").rstrip("0").rstrip(".")
    return "" if value is None else str(value)


def preview(rows):
    """打印将要粘贴的 TSV，便于和人工从 Excel 复制的效果对照。"""
    for code, debit, credit, remark in rows:
        print("  " + "\t".join(cell_text(v) for v in (code, debit, credit, remark)))
    print(f"  分录行数：{len(rows)}")


def batch_cell(path):
    """返回“批号：”右侧单元格的坐标与值。"""
    book = openpyxl.load_workbook(path, data_only=True)
    try:
        sheet = book[SHEET_NAME]
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is not None and str(cell.value).strip() == "批号：":
                    target = sheet.cell(cell.row, cell.column + 1)
                    return target.coordinate, target.value
        return None, None
    finally:
        book.close()


def main():
    parser = argparse.ArgumentParser(
        description="联调模块六：JDE 日记帐分录录入与批号回写"
    )
    parser.add_argument("--workbook", default=str(DEFAULT_WORKBOOK))
    parser.add_argument("--company", default=None,
                        help="Document Company；默认读取 房产税配置.分录公司，缺省 02501")
    parser.add_argument("--output-dir", default=str(DEFAULT_OUTPUT))
    parser.add_argument("--credential-name", default=None,
                        help="Windows 通用凭据名；默认读取 Config.xlsx")
    parser.add_argument("--headless", action="store_true",
                        help="无界面运行（默认显示浏览器）")
    parser.add_argument("--timeout", type=int, default=30000,
                        help="页面操作超时毫秒数，默认 30000")
    args = parser.parse_args()

    source = Path(args.workbook).resolve()
    if not source.is_file():
        raise SystemExit(f"计算底稿不存在：{source}")
    if args.timeout <= 0:
        raise SystemExit("--timeout 必须大于 0")

    config = Config()
    config.load(
        str(ROOT / "data" / "UserConfig.xlsx"),
        str(ROOT / "data" / "Config.xlsx"),
    )
    if args.credential_name:
        config._loader._config.setdefault("JDE配置", {})[
            "jde_credential_name"
        ] = args.credential_name
    if args.company:
        config._loader._config.setdefault("房产税配置", {})["分录公司"] = args.company

    run_dir = Path(args.output_dir).resolve() / datetime.now().strftime(
        "分录联调_%Y%m%d_%H%M%S"
    )
    run_dir.mkdir(parents=True, exist_ok=True)
    workbook = run_dir / source.name
    shutil.copy2(source, workbook)

    context = TaskContext()
    context.config = config
    flow = PropertyTaxFlow(context)

    print("=" * 60)
    print(f"源底稿  ：{source}")
    print(f"测试副本：{workbook}")
    print(f"分录公司：{flow.get('分录公司', '02501')}")
    print("=" * 60)

    print("\n[1/5] 读取计算底稿 G-J（不登录）")
    rows, description = flow.journal_rows(workbook)
    print(f"  说明：{description}")
    preview(rows)
    coordinate, previous = batch_cell(workbook)
    print(f"  副本当前批号：{SHEET_NAME}!{coordinate} = {previous}")

    print("\n[2/5] 登录 JDE")
    browser = BrowserManager(headless=args.headless)
    context.set_browser("JDE", browser)
    batch = None
    completed = False
    try:
        browser.start()
        if not JDELoginTask(context).run():
            raise RuntimeError("JDE 登录失败，无法测试分录录入")
        logger.info(f"已登录，页面：{browser.get_page().url}")

        print("\n[3/5] 录入日记帐分录（粘贴 G-J；不点取消，停在填写完成的页面）")
        batch = create_journal_entry(
            browser.get_page(),
            company=flow.get("分录公司", "02501"),
            description=description,
            gl_date=format_jde_date(date.today(), "总帐日期"),
            rows=rows,
            timeout=args.timeout,
            discard=False,
        )
        context.set("jde_journal_batch", batch)
        context.set("jde_journal_rows", len(rows))
        print(f"  JDE 批号：{batch}")

        print("\n[4/5] 校验批号回写")
        if not batch or not str(batch).strip():
            raise AssertionError("未从页面读取到批号")
        flow.write_batch_number(workbook, batch)
        coordinate, value = batch_cell(workbook)
        if coordinate is None:
            raise AssertionError(f"副本中未找到“批号：”单元格：{workbook}")
        print(f"  {SHEET_NAME}!{coordinate} = {value}")
        if str(value).strip() != str(batch).strip():
            raise AssertionError(f"批号回写不一致：页面={batch!r} 单元格={value!r}")
        if str(value).strip() == str(previous).strip():
            print("  警告：批号与测试前相同，请确认 JDE 是否真的新建了批次")
        if context.get("jde_journal_rows") != len(rows):
            raise AssertionError(
                f"录入行数不一致：{context.get('jde_journal_rows')} != {len(rows)}"
            )
        completed = True
    finally:
        if not completed:
            context.close_all_browsers()

    print("\n[5/5] 页面已填写完成，未取消、未保存")
    print(f"  批号 {batch} 已回写测试副本：{workbook}")
    print("  请在前台浏览器人工核对分录内容；核对完成后回到本窗口按回车结束。")
    try:
        input()
    except (EOFError, KeyboardInterrupt):
        pass
    context.close_all_browsers()


if __name__ == "__main__":
    main()
