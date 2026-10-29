import unittest
import openpyxl
from workflows.jde_flow.template_update import update_legacy as update_workbook


class TemplateUpdateTests(unittest.TestCase):
    def book(self):
        book = openpyxl.Workbook()
        book.remove(book.active)
        audit = book.create_sheet("房产税（从租）审核底稿")
        audit.append(["BU", "收入科目"])
        audit.append([250104, "250104.810200"])
        audit.append([None, "250104.810200"])
        audit.merge_cells("A2:A3")
        ledger = book.create_sheet("JDE明细账")
        ledger.append(["帐号", "单据 类型", "分类帐类型 1 金额", "帐户 说明"])
        ledger.append(["  250104.810200.01", "RI", -100, "租金"])
        ledger.append(["  250104.810200.01", "RI", 20, "租金"])
        ledger.append(["  250104.810200.01", "JE", -30, "租金"])
        ledger.append(["250104.8102000", "RI", -999, "不能误匹配"])
        calc = book.create_sheet("房产税（从租）计算底稿")
        calc["C8"] = .12
        calc["B9"] = 250104
        calc["C9"] = 999
        calc["B10"] = 250108
        calc["C10"] = "=LOOKUP(C4,MOH!$E$6:$AF$6,MOH!$E$7:$AF$7)"
        calc["D10"] = "=ROUND(C10*$C$8,2)"
        calc["B12"] = "合计"
        calc["G8"] = "250104.917013"
        calc["I13"] = "=-$D$12+-$I$14"
        return book

    def test_merged_bu_duplicate_accounts_signed_sum_and_dynamic_formulas(self):
        book = self.book()
        result = update_workbook(book, include_children=True)
        self.assertEqual([r["汇总金额"] for r in result], ["-80", "-30"])
        calc = book["房产税（从租）计算底稿"]
        self.assertEqual(calc["C9"].value, 80)
        self.assertEqual(calc["C10"].value, 30)
        self.assertEqual(calc["B11"].value, 250108)
        self.assertEqual(calc["C11"].value, "=LOOKUP($C$4,MOH!$E$6:$AF$6,MOH!$E$7:$AF$7)")
        self.assertEqual(calc["H8"].value, "=$D$9+$D$10")
        self.assertEqual(calc["I13"].value, "=-$D$13+-$I$14")
        self.assertEqual(calc["D13"].value, "=SUM(D9:D12)")
        self.assertEqual(calc["P7"].value, "=SUM(P5:P6)")
        book.close()

    def test_exact_match_does_not_silently_match_child_accounts(self):
        book = self.book()
        with self.assertRaisesRegex(Exception, "未匹配"):
            update_workbook(book)
        book.close()

    def test_exact_account_different_document_types_repeat_bu_rows(self):
        book = self.book()
        audit = book["房产税（从租）审核底稿"]
        audit["B2"] = "250104.810200.01"
        audit["B3"] = "250104.810200.01"
        result = update_workbook(book)
        self.assertEqual([(r["BU"], r["单据类型"], r["计税金额"]) for r in result],
                         [("250104", "RI", "80"), ("250104", "JE", "30")])
        calc = book["房产税（从租）计算底稿"]
        self.assertEqual([calc.cell(r, 2).value for r in (9, 10)], [250104, 250104])
        self.assertEqual([calc.cell(r, 3).value for r in (9, 10)], [80, 30])
        self.assertEqual(calc["H8"].value, "=$D$9+$D$10")
        book.close()

    def test_different_accounts_with_same_bu_are_not_collapsed(self):
        book = self.book()
        audit = book["房产税（从租）审核底稿"]
        audit["B2"] = "250104.810200.01"
        audit["B3"] = "250104.810100.02"
        book["JDE明细账"].append(["250104.810100.02", "RI", -50, "基本租金"])
        result = update_workbook(book)
        self.assertEqual(len(result), 3)
        calc = book["房产税（从租）计算底稿"]
        self.assertEqual([calc.cell(r, 2).value for r in (9, 10, 11)], [250104]*3)
        self.assertEqual([calc.cell(r, 3).value for r in (9, 10, 11)], [50, 80, 30])
        self.assertEqual(calc["H8"].value, "=$D$9+$D$10+$D$11")
        book.close()

    def test_fixed_bu_last_and_formulas_follow_expansion_then_shrink(self):
        book = self.book()
        audit = book["房产税（从租）审核底稿"]
        audit["B2"] = audit["B3"] = "250104.810200.01"
        ledger = book["JDE明细账"]
        ledger.append(["250104.810200.01", "AA", -10, "租金"])
        calc = book["房产税（从租）计算底稿"]
        calc["G12"] = "250108.917013"
        update_workbook(book)
        self.assertEqual(calc["B12"].value, 250108)
        self.assertEqual(calc["C12"].value, "=LOOKUP($C$4,MOH!$E$6:$AF$6,MOH!$E$7:$AF$7)")
        for r in range(9, 13):
            self.assertEqual(calc.cell(r, 4).value, f"=ROUND(C{r}*$C$8,2)")
        self.assertEqual(calc["H12"].value, "=$D$12")
        self.assertEqual(calc["D14"].value, "=SUM(D9:D13)")
        ledger.delete_rows(4, 3)
        update_workbook(book)
        self.assertEqual(calc["B10"].value, 250108)
        self.assertEqual(calc["D10"].value, "=ROUND(C10*$C$8,2)")
        self.assertEqual(calc["H12"].value, "=$D$10")
        self.assertEqual(calc["D12"].value, "=SUM(D9:D11)")
        self.assertIsNone(calc["B14"].value)
        book.close()

    def test_fixed_bu_missing_is_created_with_formula(self):
        book = self.book()
        calc = book["房产税（从租）计算底稿"]
        calc["B10"] = calc["C10"] = calc["D10"] = None
        update_workbook(book, include_children=True)
        self.assertEqual(calc["B11"].value, 250108)
        self.assertEqual(calc["C11"].value, "=LOOKUP($C$4,MOH!$E$6:$AF$6,MOH!$E$7:$AF$7)")
        self.assertEqual(calc["D11"].value, "=ROUND(C11*$C$8,2)")
        book.close()
