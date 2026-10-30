"""本地临时工作簿验证，不连接 JDE。"""
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
import json
from unittest.mock import MagicMock, patch

import openpyxl

from config.config_loader import ConfigError
from core.context import TaskContext
from workflows.jde_flow.property_tax import PropertyTaxFlow, HEADER_DEFINITION, PendingModule


class PropertyTaxTests(unittest.TestCase):
    def setUp(self):
        self.tmp = TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.values = {"公司": "001", "会计期间": "2026-10", "明细账文件": "ledger.xlsx",
                       "重算引擎": "openpyxl"}
        context = TaskContext()
        class Config:
            def get(_, key, default=None):
                return self.values.get(key.split(".", 1)[1], default)
            def get_path(_, kind):
                return self.root
        context.config = Config()
        self.flow = PropertyTaxFlow(context)
        def pending_update(path):
            raise PendingModule("模块四：此测试仅验证模块三")
        self.flow.update_template = pending_update
        self.headers = [c["template_header"] for c in json.loads(
            HEADER_DEFINITION.read_text(encoding="utf-8"))["columns"]]
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = "JDE明细账"
        sheet.append(self.headers)
        sheet.append(["旧记录"] * 46)
        sheet.append(["旧尾行"] * 46)
        sheet["H2"].number_format = "#,##0.00"
        book.create_sheet("计算")["A1"] = "=SUM('JDE明细账'!A2:A9)"
        book.save(self.root / "房产税计税表模板.xlsx")
        book.close()
        book = openpyxl.Workbook()
        # 源列序与模板相反。
        book.active.append(list(reversed(self.headers)))
        book.active.append([123.45 if title == "分类帐类型 1 金额" else
                            "00100" if title == "帐号" else None
                            for title in reversed(self.headers)])
        book.save(self.root / "ledger.xlsx")
        book.close()

    def test_reorders_columns_preserves_formulas_and_stops_before_comparison(self):
        self.flow.export_ledger = lambda: self.root / "ledger.xlsx"
        self.flow.export_r55gl1005 = lambda: None
        self.flow.run()
        context = self.flow.context
        self.assertEqual(context.get("property_tax_status"), "pending")
        self.assertIn("模块四", context.get("property_tax_pending_reason"))
        self.assertIsNone(context.get("journal_required"))
        path = context.get("property_tax_path")
        self.assertEqual(path.name, "房产税计税表 001 202610.xlsx")
        book = openpyxl.load_workbook(path)
        try:
            self.assertEqual(book["JDE明细账"]["H2"].value, 123.45)
            self.assertEqual(book["JDE明细账"]["A2"].value, "00100")
            self.assertTrue(all(c.value is None for c in book["JDE明细账"][3]))
            self.assertEqual(book["JDE明细账"]["H2"].number_format, "#,##0.00")
            self.assertEqual(book["计算"]["A1"].value, "=SUM('JDE明细账'!A2:A9)")
        finally:
            book.close()
        with self.assertRaises(ConfigError):
            self.flow.create_template(self.root / "ledger.xlsx")

    def test_missing_export_company_rejected(self):
        # 明细账导出始终执行：缺少自公司/至公司时明确报错，不再静默跳过。
        with self.assertRaisesRegex(ConfigError, "模块二缺少导出参数"):
            self.flow.run()
        self.assertEqual(self.flow.context.get("property_tax_status"), "failed")

    def test_incomplete_mapping_rejected_without_output(self):
        self.values["标题映射"] = '{"不存在的科目": "帐号"}'
        book = openpyxl.load_workbook(self.root / "ledger.xlsx")
        book.active.cell(1, 1).value = "未知标题"
        book.save(self.root / "ledger.xlsx")
        book.close()
        with self.assertRaises(ConfigError):
            self.flow.create_template(self.root / "ledger.xlsx")
        self.assertIsNone(self.flow.context.get("property_tax_path"))

    def test_partial_alias_mapping_keeps_other_identical_headers(self):
        book = openpyxl.load_workbook(self.root / "ledger.xlsx")
        book.active.cell(1, 46).value = "源账号"
        book.save(self.root / "ledger.xlsx")
        book.close()
        self.values["标题映射"] = '{"源账号": "帐号"}'
        self.flow.export_ledger = lambda: self.root / "ledger.xlsx"
        self.flow.export_r55gl1005 = lambda: None
        self.flow.run()
        self.assertEqual(self.flow.context.get("jde_ledger_mapping")["源账号"], "帐号")

    def test_missing_source_column_rejected(self):
        book = openpyxl.load_workbook(self.root / "ledger.xlsx")
        book.active.delete_cols(1)
        book.save(self.root / "ledger.xlsx")
        book.close()
        with self.assertRaisesRegex(ConfigError, "缺少模板所需列"):
            self.flow.create_template(self.root / "ledger.xlsx")

    @patch("workflows.jde_flow.property_tax.convert_legacy_excel_to_xlsx")
    @patch("workflows.jde_flow.property_tax.download_r55gl1005")
    def test_r55_export_saves_raw_xls_and_converts_to_configured_xlsx(
        self, download, convert
    ):
        page = MagicMock()
        browser = MagicMock()
        browser.get_page.return_value = page
        self.flow.context.set_browser("JDE", browser)
        self.values["明细账自公司"] = "02501"
        self.values["R55GL1005文件"] = "本期报表.xls"
        raw_path = self.root / "本期报表.xls"
        download.return_value = raw_path

        result = self.flow.export_r55gl1005()

        expected_xlsx = self.root / "本期报表.xlsx"
        download.assert_called_once_with(
            page,
            company="02501",
            destination=raw_path,
        )
        convert.assert_called_once_with(raw_path, expected_xlsx)
        self.assertEqual(result, expected_xlsx)
        self.assertEqual(
            self.flow.context.get("jde_r55gl1005_path"), expected_xlsx
        )

    def test_run_exports_r55_after_ledger_before_template_creation(self):
        events = []
        ledger_path = self.root / "ledger.xlsx"
        template_path = self.root / "template.xlsx"
        self.flow.export_ledger = lambda: events.append("ledger") or ledger_path
        self.flow.export_r55gl1005 = lambda: events.append("r55")
        self.flow.create_template = lambda path: (
            events.append(("template", path)) or template_path
        )
        self.flow.update_template = lambda path: events.append(("update", path))

        def stop_before_comparison(path):
            events.append(("compare", path))
            raise PendingModule("stop after verifying export order")

        self.flow.compare_values = stop_before_comparison
        self.flow.run()

        self.assertEqual(
            events,
            [
                "ledger",
                "r55",
                ("template", ledger_path),
                ("update", template_path),
                ("compare", template_path),
            ],
        )
