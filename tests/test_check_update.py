import unittest
from pathlib import Path
import openpyxl
from config.config_loader import ConfigError
from workflows.jde_flow.template_update import update_workbook
from workflows.jde_flow.audit_update import update_audit


class CheckUpdateTests(unittest.TestCase):
    def book(self):
        path = Path(__file__).resolve().parents[1] / 'data/房产税计税表 太古汇 202609更新.xlsx'
        b = openpyxl.load_workbook(path)
        self.addCleanup(b.close)
        return b

    def test_real_template_report_and_formulas(self):
        b = self.book()
        records = update_workbook(b)
        self.assertEqual(len(records), 11)
        s = b['房产税（从租）计算底稿']
        self.assertEqual(s['C9'].value, '=-SUM(P5)')
        self.assertEqual(s['C10'].value, '=-SUM(P6)')
        self.assertEqual(s['C11'].value, '=-SUM(P7,P8,P9,P10)')
        self.assertEqual(s['M6'].value, '非JE')
        self.assertIn('<>MID(M6,2,99)', s['P6'].value)
        self.assertEqual(s['H8'].value, '=$D$9+$D$10+$D$11')
        audit = b['房产税（从租）审核底稿']
        self.assertEqual(audit['F31'].value, '其他业务收入')
        self.assertEqual(audit['N14'].value, "=N11+N12-'房产税（从租）计算底稿'!C17")
        self.assertIn('$I$9', b.defined_names['蓝色'].attr_text)
        update_workbook(b)
        self.assertEqual(s['I13'].value, '=-$D$17+-$I$14')

    def test_signed_filters_blank_types_and_unmatched_account(self):
        b = self.book()
        l = b['JDE明细账']
        l.delete_rows(2,l.max_row)
        cols = {c.value:c.column for c in l[1]}
        for doc, amount in [('JE',-30),('RI',-80),(None,20)]:
            r = l.max_row+1
            l.cell(r,cols['帐号']).value = ' 250104.810200.01'
            l.cell(r,cols['单据 类型']).value = doc
            l.cell(r,cols['分类帐类型 1 金额']).value = amount
        records = update_workbook(b)
        self.assertEqual([r['汇总金额'] for r in records[:3]], ['-30','-60','0'])

    def test_duplicate_and_unclassified_accounts_block(self):
        b = self.book()
        b['房产税（从租）计算底稿']['M6'] = 'JE'
        with self.assertRaisesRegex(ConfigError,'重复'):
            update_workbook(b)
        s = b['R55GL1005']
        r = next(r for r in range(6,s.max_row+1) if s.cell(r,5).value == '主营业务收入')
        s.cell(r+1,5).value = 'NEW ACCOUNT'
        with self.assertRaisesRegex(ConfigError,'未分类'):
            update_audit(b,16)
