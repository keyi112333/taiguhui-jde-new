"""离线联调：用已导出的文件跑“写入模板 → 更新底稿 → 重算并比对金额”。

不登录 JDE、不下载文件，直接使用 data/JDE流程 下的模板和两次导出结果；
R55GL1005 若给的是 .xls，会先按线上流程转成 .xlsx（优先复用同名 .xlsx）。

用法：
    python tests/manual_template_pipeline.py
    python tests/manual_template_pipeline.py --ledger "data/JDE流程/03_输出/JDE明细账 20261029_092539.xlsx"
"""
import argparse
from datetime import datetime
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import openpyxl

from core.context import TaskContext
from config.config_loader import ConfigError
from log.loggingConfig import logger
from workflows.jde_flow.property_tax import PendingModule, PropertyTaxFlow
from workflows.jde_flow.template_check import convert_legacy_excel_to_xlsx

DEFAULT_DIR = ROOT / "data/JDE流程"
DEFAULT_TEMPLATE = DEFAULT_DIR / "02_输入/房产税计税表 太古汇 202609更新.xlsx"
DEFAULT_LEDGER = DEFAULT_DIR / "03_输出/JDE明细账 20261029_092052.xlsx"
DEFAULT_REPORT = DEFAULT_DIR / "03_输出/R55GL1005_20261029_092155.xls"


class LocalConfig:
    """只提供本次联调所需的配置读取，不依赖 Config.xlsx。"""

    def __init__(self, values, input_dir, output_dir):
        self.values = values
        self.input_dir = input_dir
        self.output_dir = output_dir

    def get(self, key, default=None):
        return self.values.get(key, default)

    def get_path(self, kind):
        return self.output_dir if kind == "output" else self.input_dir


def xlsx_report(report, work_dir, force_convert=False):
    """按线上规则把 R55GL1005 转成 openpyxl 可读的 .xlsx。"""
    report = report.resolve()
    if report.suffix.lower() == ".xlsx":
        return report
    if not report.is_file():
        raise FileNotFoundError(f"R55GL1005 文件不存在：{report}")
    target = report.with_suffix(".xlsx")
    if target.is_file() and not force_convert:
        logger.info(f"复用已转换的 R55GL1005：{target}")
        return target
    target = work_dir / report.with_suffix(".xlsx").name
    logger.info(f"转换 R55GL1005：{report} -> {target}")
    convert_legacy_excel_to_xlsx(report, target)
    return target


def check_formulas(path):
    """更新后、重算前检查关键公式仍在，避免只比对缓存值。"""
    book = openpyxl.load_workbook(path, data_only=False)
    try:
        audit = book["房产税（从租）审核底稿"]
        rows = [r for r in range(8, audit.max_row + 1) if audit.cell(r, 6).value is not None]
        if not rows:
            raise AssertionError("审核底稿未生成科目行")
        for r in rows:
            value = audit.cell(r, 9).value
            if value != f"=H{r}-G{r}":
                raise AssertionError(f"审核底稿 I{r} 公式不完整：{value!r}")
        if "蓝色" not in book.defined_names:
            raise AssertionError("缺少蓝色区域命名范围")
        return {
            "审核科目行": len(rows),
            "蓝色范围": book.defined_names["蓝色"].attr_text,
            "审核差额 N14": audit["N14"].value,
            "计税收入 N11": audit["N11"].value,
            "MOH N12": audit["N12"].value,
        }
    finally:
        book.close()


def main():
    parser = argparse.ArgumentParser(description="离线联调导出表写入模板、内容更新与金额比对")
    parser.add_argument("--template", default=str(DEFAULT_TEMPLATE))
    parser.add_argument("--ledger", default=str(DEFAULT_LEDGER))
    parser.add_argument("--report", default=str(DEFAULT_REPORT))
    parser.add_argument("--company", default="太古汇")
    parser.add_argument("--period", default="202609")
    parser.add_argument("--tolerance", default="0.01",
                        help="金额比对容差；Excel 浮点求和会有 1E-10 级误差，默认 0.01")
    parser.add_argument("--output-dir", default=str(DEFAULT_DIR / "03_输出"))
    parser.add_argument("--force-convert", action="store_true",
                        help="忽略同名 .xlsx，重新从 .xls 转换 R55GL1005")
    args = parser.parse_args()

    template = Path(args.template).resolve()
    ledger = Path(args.ledger).resolve()
    for path in (template, ledger):
        if not path.is_file():
            raise SystemExit(f"输入文件不存在：{path}")

    run_dir = Path(args.output_dir).resolve() / datetime.now().strftime("模板联调_%Y%m%d_%H%M%S")
    run_dir.mkdir(parents=True, exist_ok=True)
    report = xlsx_report(Path(args.report), run_dir, args.force_convert)

    values = {
        "房产税配置.模板文件": str(template),
        "房产税配置.R55GL1005文件": str(report),
        "房产税配置.公司": args.company,
        "房产税配置.会计期间": args.period,
        "房产税配置.输出目录": str(run_dir),
        "房产税配置.明细账标题行": 1,
        "房产税配置.核对容差": args.tolerance,
    }
    context = TaskContext()
    context.config = LocalConfig(values, template.parent, run_dir)
    flow = PropertyTaxFlow(context)

    print("=" * 60)
    print(f"模板  ：{template}")
    print(f"明细账：{ledger}")
    print(f"报表  ：{report}")
    print(f"输出  ：{run_dir}")
    print(f"容差  ：{args.tolerance}")
    print("=" * 60)

    print("\n[1/3] 导出表写入模板（create_template）")
    path = flow.create_template(ledger)
    print(f"  生成文件：{path}")
    print(f"  写入明细行：{context.get('jde_ledger_rows')}")
    if not Path(path).is_file() or Path(path).stat().st_size == 0:
        raise SystemExit("模板生成失败：输出文件不存在或为空")

    print("\n[2/3] 模板内容更新（update_template）")
    flow.update_template(path)
    totals = context.get("property_tax_bu_totals") or []
    print(f"  Check 汇总条数：{len(totals)}")
    for item in check_formulas(path).items():
        print(f"  {item[0]}：{item[1]}")
    print(f"  更新后文件：{path}（{Path(path).stat().st_size} 字节）")

    print("\n[3/3] 重算并比对金额（compare_values）")
    failure = None
    try:
        flow.compare_values(path)
    except PendingModule as exc:
        if context.get("property_tax_check_status") != "passed":
            failure = exc
        else:
            print(f"  提示：{exc}")
    except ConfigError as exc:
        failure = exc

    status = context.get("property_tax_check_status")
    results = context.get("property_tax_check_results") or []
    for item in results:
        flag = "通过" if item["通过"] else "不通过"
        print(f"  [{flag}] {item['核对项']}：左={item['左值']} 右={item['右值']} "
              f"差额={item['差额']} 容差={item['容差']}")

    print(f"\n金额比对状态：{status}")
    if status != "passed":
        print(f"失败原因：{failure}")
        raise SystemExit(1)
    print("联调完成：模板写入、内容更新、重算比对均成功。")


if __name__ == "__main__":
    main()
