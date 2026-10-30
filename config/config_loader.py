#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
配置读取模块
"""
from pathlib import Path
from typing import Any, Dict, Optional, Union

try:
    import openpyxl
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


class ConfigError(Exception):
    """配置错误"""
    pass


class ConfigLoader:
    """配置加载器"""

    def __init__(self, config_path: Optional[Union[str, Path]] = None):
        self._config: Dict[str, Any] = {}
        self._rows: Dict[str, list] = {}
        if config_path:
            self.load(config_path)

    def load(self, config_path: Union[str, Path]) -> Dict[str, Any]:
        """加载Excel配置文件"""
        if not HAS_OPENPYXL:
            raise ConfigError("需要安装 openpyxl: pip install openpyxl")

        path = Path(config_path)
        if not path.exists():
            raise ConfigError(f"配置文件不存在: {config_path}")

        workbook = openpyxl.load_workbook(path, data_only=True)

        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            self._config[sheet_name] = self._read_sheet(sheet)
            self._rows[sheet_name] = self._read_rows(sheet)

        workbook.close()
        return self._config

    def _read_sheet(self, sheet) -> Dict[str, Any]:
        """读取工作表为字典"""
        result = {}
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if len(row) >= 2 and row[0] and row[1] is not None:
                result[row[0]] = row[1]
        return result

    def _read_rows(self, sheet) -> list:
        """按首行标题把工作表读成多行字典，空行跳过（用于允许多行的配置页）。"""
        header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        titles = [str(value).strip() if value is not None else "" for value in header_row]
        rows = []
        for values in sheet.iter_rows(min_row=2, values_only=True):
            if all(value is None or not str(value).strip() for value in values):
                continue
            row = {}
            for index, title in enumerate(titles):
                if not title:
                    continue
                value = values[index] if index < len(values) else None
                row[title] = value.strip() if isinstance(value, str) else value
            rows.append(row)
        return rows

    def get_rows(self, sheet_name: str) -> list:
        """获取多行配置；未配置该页时返回空列表。"""
        return [dict(row) for row in self._rows.get(sheet_name, [])]

    def set(self, key: str, value: Any) -> None:
        """覆盖某个配置项（用于按账号切换运行参数）。"""
        if '.' not in key:
            raise ConfigError(f"配置键需为 sheet.配置项 格式：{key}")
        sheet, k = key.split('.', 1)
        self._config.setdefault(sheet, {})[k] = value

    def get(self, key: str, default: Any = None) -> Any:
        """获取配置项（支持 sheet.key 格式）"""
        if '.' in key:
            sheet, k = key.split('.', 1)
            return self._config.get(sheet, {}).get(k, default)
        return self._config.get(key, default)

    @property
    def config(self) -> Dict[str, Any]:
        return self._config.copy()
